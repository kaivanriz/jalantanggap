"""Generate a deterministic LAKSA-format synthetic complaint workbook for JalanTanggap Fase 1."""
import csv
import math
import random
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "sample-tangerang"
OUT_XLSX = DATA / "laporan_laksa_sintetis_fase1.xlsx"
OUT_TRUTH = DATA / "ground_truth_laksa_sintetis_fase1.csv"

ROAD_HASHTAGS = ["#JALAN RUSAK", "#JALAN BERLUBANG", "#PERBAIKAN JALAN", "#ASPAL RUSAK"]
ROAD_ISSUES = [
    ("lubang", "Ada lubang besar di {jalan} dekat {landmark}, pengendara sering menghindar mendadak."),
    ("retak", "Aspal {jalan} dekat {landmark} retak memanjang dan makin tidak rata."),
    ("aspal_mengelupas", "Lapisan aspal {jalan} sekitar {landmark} mengelupas dan kerikil berserakan."),
    ("permukaan_bergelombang", "{jalan} dekat {landmark} bergelombang sehingga kendaraan harus melambat."),
]
NOISE = [
    ("#DRAINASE BARU", "drainase", "Mohon pembangunan drainase baru di sekitar {landmark}, saat hujan air sering menggenang."),
    ("#TIANG NON TELKOM PLN", "utilitas_tiang", "Ada tiang utilitas miring dekat {landmark}, mohon ditindaklanjuti."),
    ("#PJU", "pju", "Lampu penerangan jalan dekat {landmark} mati dan malam hari gelap."),
    ("#POHON", "pohon", "Ada dahan pohon besar menjulur ke badan jalan dekat {landmark}."),
    ("#TROTOAR", "trotoar", "Trotoar dekat {landmark} rusak dan mengganggu pejalan kaki."),
]
LANDMARKS = ["sekolah", "pasar", "puskesmas", "halte", "persimpangan", "pertokoan"]
SOURCES = ["LAKSA", "TIKTOK", "INSTAGRAM", "WHATSAPP", "WEB"]
NAMES = ["warga.tangerang", "rudi", "siti", "andi", "anonim", "fitri", "budi"]


def load_network():
    with (DATA / "nodes.csv").open(encoding="utf-8") as f:
        nodes = {r["id_node"]: r for r in csv.DictReader(f)}
    with (DATA / "segmen.csv").open(encoding="utf-8") as f:
        segs = list(csv.DictReader(f))
    return nodes, segs


def offset_midpoint(a, b, seed):
    rng = random.Random(seed)
    lat1, lon1 = float(a["latitude"]), float(a["longitude"])
    lat2, lon2 = float(b["latitude"]), float(b["longitude"])
    t = rng.uniform(0.30, 0.72)
    lat = lat1 + (lat2 - lat1) * t
    lon = lon1 + (lon2 - lon1) * t
    dy = (lat2 - lat1) * 111320
    dx = (lon2 - lon1) * 111320 * math.cos(math.radians(lat))
    length = math.hypot(dx, dy) or 1
    ux, uy = dx / length, dy / length
    cross = rng.uniform(-18, 18)
    xoff, yoff = -uy * cross, ux * cross
    lat += yoff / 111320
    lon += xoff / (111320 * math.cos(math.radians(lat)))
    return round(lat, 9), round(lon, 9)


def main():
    DATA.mkdir(parents=True, exist_ok=True)
    nodes, segs = load_network()
    road_segs = segs[:20]

    rows = []
    truth = []
    start = datetime(2026, 8, 1, 7, 0)
    rid = 151000

    for sidx, seg in enumerate(road_segs):
        a, b = nodes[seg["from_node"]], nodes[seg["to_node"]]
        issue, base_text = ROAD_ISSUES[sidx % len(ROAD_ISSUES)]
        landmark = LANDMARKS[sidx % len(LANDMARKS)]
        group = f"DG-{seg['id_segmen']}-01"
        for j in range(4):
            rid += 1
            lat, lon = offset_midpoint(a, b, f"{seg['id_segmen']}-{j}")
            text = base_text.format(jalan=seg["nama"], landmark=landmark)
            if j == 1:
                text = "Mohon dicek, " + text[0].lower() + text[1:] + " Sudah beberapa hari."
            elif j == 2:
                text = text.replace("dekat", "di sekitar") + " Kondisinya mengganggu saat ramai."
            elif j == 3:
                text = "Lapor lagi: " + text + " Belum tertangani."
            t = start + timedelta(days=(sidx * 2 + j) % 38, hours=(sidx + j) % 11, minutes=(sidx * 9 + j * 17) % 60)
            source = SOURCES[(sidx + j) % len(SOURCES)]
            rows.append([
                len(rows) + 1, str(rid), NAMES[(sidx + j) % len(NAMES)], t,
                ROAD_HASHTAGS[sidx % len(ROAD_HASHTAGS)], "DPUPR", text,
                f"{seg['nama']}, dekat {landmark} simulasi", str(lat).replace(".", ","), str(lon).replace(".", ","),
                "Menunggu" if j % 3 else "Sedang Proses", "", "", "", "", source, "Aduan", "", ""
            ])
            truth.append({
                "id_pengaduan": str(rid), "expected_id_segmen": seg["id_segmen"],
                "expected_category": "jalan_rusak", "expected_indication": issue,
                "expected_road_related": True, "expected_duplicate_group": group,
                "expected_is_duplicate": j > 0, "is_simulation": True
            })

    for i in range(20):
        rid += 1
        seg = road_segs[i % len(road_segs)]
        a, b = nodes[seg["from_node"]], nodes[seg["to_node"]]
        lat, lon = offset_midpoint(a, b, f"noise-{i}")
        tag, category, tpl = NOISE[i % len(NOISE)]
        lm = LANDMARKS[(i + 2) % len(LANDMARKS)] + " simulasi"
        t = start + timedelta(days=(i * 3) % 38, hours=(i + 2) % 10, minutes=(i * 11) % 60)
        text = tpl.format(landmark=lm)
        rows.append([
            len(rows) + 1, str(rid), NAMES[i % len(NAMES)], t, tag, "DPUPR", text,
            f"{seg['nama']}, dekat {lm}", str(lat).replace(".", ","), str(lon).replace(".", ","),
            "Menunggu", "", "", "", "", SOURCES[i % len(SOURCES)], "Aduan", "", ""
        ])
        truth.append({
            "id_pengaduan": str(rid), "expected_id_segmen": "",
            "expected_category": category, "expected_indication": category,
            "expected_road_related": False, "expected_duplicate_group": "",
            "expected_is_duplicate": False, "is_simulation": True
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Pengaduan"
    ws.merge_cells("A1:S2")
    ws["A1"] = "LAPORAN PENGADUAN LAKSA — DATA SINTETIS FASE 1"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["No","ID Pengaduan","Dari","Tgl.Laporan","Hastag","Tujuan","Isi Pengaduan","Lokasi","Lat","Lng","Status Lapor","Tgl Proses","Waktu Respon","Tgl Selesai","Waktu TL","Sumber","Jenis Aduan","Jawaban",""]
    for c, v in enumerate(headers, 1):
        ws.cell(6, c, v)
    ws.cell(7, 18, "Tanggal")
    ws.cell(7, 19, "Jawaban")
    ws.merge_cells("R6:S6")

    thin = Side(style="thin", color="808080")
    fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws.iter_rows(min_row=6, max_row=7, min_col=1, max_col=19):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, 8):
        for c_idx, val in enumerate(row, 1):
            ws.cell(r_idx, c_idx, val)
        ws.cell(r_idx, 4).number_format = "yyyy-mm-dd hh:mm:ss"
        for c in range(1, 20):
            ws.cell(r_idx, c).alignment = Alignment(vertical="top", wrap_text=True)

    widths = [7,14,20,20,24,12,58,45,16,16,16,20,18,20,14,12,14,20,60]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A8"
    wb.save(OUT_XLSX)

    with OUT_TRUTH.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(truth[0]))
        writer.writeheader()
        writer.writerows(truth)

    print(f"Wrote {OUT_XLSX}")
    print(f"Wrote {OUT_TRUTH}")
    print(f"Rows: {len(rows)}")


if __name__ == "__main__":
    main()
